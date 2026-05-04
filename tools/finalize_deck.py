"""Apply final edits to the Sourcing Agent Gamma export.

Reads:
  assets/deck/sourcing-agent-original.pptx           (untouched)
  assets/deck-diagrams/exports/slide-*.png           (high-res diagram PNGs)

Writes:
  assets/deck/sourcing-agent-final.pptx

Every edit is expressed as a small helper call so the intent is readable. Image
swaps preserve aspect ratio by sizing to a target width and computing the
corresponding height from the source PNG.
"""

from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from PIL import Image


REPO = Path(__file__).resolve().parent.parent
SRC = REPO / "assets/deck/sourcing-agent-original.pptx"
DST = REPO / "assets/deck/sourcing-agent-final.pptx"
IMG = REPO / "assets/deck-diagrams/exports"


# ---------- helpers ------------------------------------------------------------


def png_aspect(path: Path) -> float:
    """Return width / height for the given PNG."""
    with Image.open(path) as im:
        return im.width / im.height


def size_for_width(path: Path, width_in: float) -> tuple[float, float]:
    """Compute (width_in, height_in) that preserves aspect at the given width."""
    return width_in, width_in / png_aspect(path)


def size_for_height(path: Path, height_in: float) -> tuple[float, float]:
    """Compute (width_in, height_in) that preserves aspect at the given height."""
    return height_in * png_aspect(path), height_in


def center_x(slide_width_in: float, width_in: float) -> float:
    return (slide_width_in - width_in) / 2


def remove_shape(shape) -> None:
    """Remove a shape from its parent spTree."""
    sp = shape._element
    sp.getparent().remove(sp)


def find_shape_by_text(slide, needle: str):
    """Return the first shape whose text frame contains `needle`, or None."""
    for shape in slide.shapes:
        if shape.has_text_frame and needle in shape.text_frame.text:
            return shape
    return None


def find_shape_by_index(slide, idx: int):
    return list(slide.shapes)[idx]


def replace_image_at(slide, old_image_shape, png_path: Path, left_in: float,
                     top_in: float, width_in: float, height_in: float) -> None:
    """Remove an existing picture and insert a fresh one at the given rect."""
    remove_shape(old_image_shape)
    slide.shapes.add_picture(
        str(png_path),
        Inches(left_in),
        Inches(top_in),
        width=Inches(width_in),
        height=Inches(height_in),
    )


def set_text_preserve_formatting(text_frame, new_text: str) -> None:
    """Replace the text of a text frame, preserving the formatting of its first run."""
    # Capture formatting from the first run if one exists.
    first_p = text_frame.paragraphs[0]
    template_run = first_p.runs[0] if first_p.runs else None

    # Wipe the text frame: clear all paragraphs except the first, clear the first.
    for p in list(text_frame.paragraphs[1:]):
        p._p.getparent().remove(p._p)
    p0 = text_frame.paragraphs[0]
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)

    # Write new text, one paragraph per newline.
    lines = new_text.split("\n")
    # First line goes into the existing paragraph.
    run = p0.add_run()
    run.text = lines[0]
    if template_run is not None:
        copy_run_formatting(template_run, run)
    # Subsequent lines become new paragraphs.
    for line in lines[1:]:
        p = text_frame.add_paragraph()
        r = p.add_run()
        r.text = line
        if template_run is not None:
            copy_run_formatting(template_run, r)


def copy_run_formatting(src_run, dst_run) -> None:
    """Copy font formatting (name, size, bold, italic, color) from src to dst."""
    src_font = src_run.font
    dst_font = dst_run.font
    if src_font.name:
        dst_font.name = src_font.name
    if src_font.size:
        dst_font.size = src_font.size
    if src_font.bold is not None:
        dst_font.bold = src_font.bold
    if src_font.italic is not None:
        dst_font.italic = src_font.italic
    # Color: only copy if the source has an explicit rgb color.
    try:
        if src_font.color and src_font.color.rgb:
            dst_font.color.rgb = src_font.color.rgb
    except Exception:
        pass


# ---------- main ---------------------------------------------------------------


def main() -> None:
    prs = Presentation(str(SRC))
    slide_w_in = prs.slide_width / 914400
    slide_h_in = prs.slide_height / 914400
    slides = list(prs.slides)
    assert len(slides) == 9, f"Expected 9 slides, got {len(slides)}"

    # --- Slide 1: remove the thin horizontal rule under the subtitle ---
    # The rule is the last text shape ("Shape 4"), height ~0.02in, empty text.
    s1 = slides[0]
    for shape in list(s1.shapes):
        if (shape.has_text_frame and shape.text_frame.text == ""
                and shape.height is not None and shape.height < Emu(50000)):
            remove_shape(shape)

    # --- Slide 2: swap hub-and-spoke diagram, remove redundant bullet list ---
    s2 = slides[1]
    # Find the existing picture shape.
    old_pic = next(sh for sh in s2.shapes if sh.shape_type == 13)
    png = IMG / "slide-02-hub-and-spoke.png"
    # Position: right column, centered vertically under the eyebrow/title.
    # Original picture was 6.19" x 3.48" at (8.63, 3.17). The new SVG aspect is
    # 0.909 (nearly square). Give it a roomy box.
    new_h = 5.4
    new_w = new_h * png_aspect(png)  # ~ 4.91"
    new_left = slide_w_in - new_w - 0.9  # 0.9" right margin
    new_top = (slide_h_in - new_h) / 2 + 0.3
    replace_image_at(s2, old_pic, png, new_left, new_top, new_w, new_h)
    # Remove the bulleted list shape (Text 4): diagram now carries it.
    bullet_shape = find_shape_by_text(s2, "search strategy")
    if bullet_shape is not None:
        remove_shape(bullet_shape)

    # --- Slide 3: swap the tiny horizontal flow for a proper full-width banner ---
    s3 = slides[2]
    old_pic = next(sh for sh in s3.shapes if sh.shape_type == 13)
    png = IMG / "slide-03-horizontal-flow.png"
    # Banner aspect is 9.06. Make it ~14" wide => ~1.55" tall.
    new_w = 14.0
    new_h = new_w / png_aspect(png)  # ~ 1.55"
    new_left = center_x(slide_w_in, new_w)
    # Position below the 6-step grid. Bottom of step 6 row is at y ~ 6.86in
    # (step body at 6.57in + 0.29in). Give it some breathing room.
    new_top = 7.25
    # Make sure it doesn't hang off the slide.
    if new_top + new_h > slide_h_in - 0.1:
        new_top = slide_h_in - new_h - 0.1
    replace_image_at(s3, old_pic, png, new_left, new_top, new_w, new_h)

    # --- Slide 5: swap LinkedIn diagram; remove the stray second image ---
    s5 = slides[4]
    pics = [sh for sh in s5.shapes if sh.shape_type == 13]
    # Keep the larger one (the real diagram), remove the second (decorative layer).
    pics.sort(key=lambda p: p.image and len(p.image.blob), reverse=True)
    main_pic, stray_pic = pics[0], pics[1]
    remove_shape(stray_pic)
    png = IMG / "slide-05-search-families.png"
    # Aspect 1.146 (slightly wider than tall). Target a ~5.6in tall box on the right.
    new_h = 5.6
    new_w = new_h * png_aspect(png)  # ~ 6.42"
    new_left = slide_w_in - new_w - 0.6
    new_top = 1.9
    replace_image_at(s5, main_pic, png, new_left, new_top, new_w, new_h)
    # Fold orphan summary sentence into the main body paragraph.
    # Current: Text 2 ends with "... consistent, reviewable pace." and there's
    # a dangling line later. In the Gamma export the summary was actually folded
    # into Text 2 already (ends with " Every action is traceable..."). Verify
    # and, if not, rewrite Text 2.
    body_shape = find_shape_by_index(s5, 2)
    if body_shape.has_text_frame and "Every action is traceable" not in body_shape.text_frame.text:
        new_body = (
            "Sourcing Agent works inside LinkedIn Recruiter the way a strong "
            "sourcer would, but at a consistent, reviewable pace. Every action "
            "is traceable, and every search change is bounded and reviewable."
        )
        set_text_preserve_formatting(body_shape.text_frame, new_body)

    # --- Slide 6: swap GitHub diagram; remove stray second image ---
    s6 = slides[5]
    pics = [sh for sh in s6.shapes if sh.shape_type == 13]
    pics.sort(key=lambda p: p.image and len(p.image.blob), reverse=True)
    main_pic, stray_pic = pics[0], pics[1]
    remove_shape(stray_pic)
    png = IMG / "slide-06-seed-expansion.png"
    # Aspect 1.241. Target ~5.6" tall.
    new_h = 5.6
    new_w = new_h * png_aspect(png)  # ~ 6.95"
    new_left = slide_w_in - new_w - 0.6
    new_top = 1.9
    replace_image_at(s6, main_pic, png, new_left, new_top, new_w, new_h)

    # --- Slide 7: nuke pinwheel + floating dots + floating labels, insert loop ---
    s7 = slides[6]
    # Everything except the first 8 text shapes (eyebrow, title, intro sentence,
    # "From a single run:", its body, "Across many runs:", its body, closing sentence)
    # is pinwheel debris. Those 8 shapes are indices 0..7. Everything from index 8
    # onward is pinwheel-related (the illustration, 4 dots, 4 labels).
    keep_first_n = 8
    all_shapes = list(s7.shapes)
    for shape in all_shapes[keep_first_n:]:
        remove_shape(shape)
    png = IMG / "slide-07-loop.png"
    # Aspect 3.074. Target ~12" wide => ~3.9" tall.
    new_w = 12.0
    new_h = new_w / png_aspect(png)  # ~ 3.9"
    new_left = center_x(slide_w_in, new_w)
    new_top = 4.3
    if new_top + new_h > slide_h_in - 0.3:
        new_top = slide_h_in - new_h - 0.3
    s7.shapes.add_picture(
        str(png),
        Inches(new_left),
        Inches(new_top),
        width=Inches(new_w),
        height=Inches(new_h),
    )

    # --- Slide 8: swap governed-fleet diagram; remove stray second image; rewrite paragraph 2 ---
    s8 = slides[7]
    pics = [sh for sh in s8.shapes if sh.shape_type == 13]
    pics.sort(key=lambda p: p.image and len(p.image.blob), reverse=True)
    main_pic, stray_pic = pics[0], pics[1]
    remove_shape(stray_pic)
    png = IMG / "slide-08-governed-fleet.png"
    # Aspect 1.8 (clearly wider than tall). Target a box that sits well on the right.
    new_h = 5.0
    new_w = new_h * png_aspect(png)  # ~ 9.0"
    # If too wide to fit to the right of the body column (body ends at ~7.76in),
    # shrink it.
    max_w = slide_w_in - 8.2 - 0.3  # body copy column ends ~7.76, allow a small gap
    if new_w > max_w:
        new_w = max_w
        new_h = new_w / png_aspect(png)
    new_left = slide_w_in - new_w - 0.3
    new_top = 2.4
    replace_image_at(s8, main_pic, png, new_left, new_top, new_w, new_h)
    # Rewrite paragraph 2 (Text 3: "The next step is orchestration: ...").
    p2_shape = find_shape_by_text(s8, "The next step is orchestration")
    if p2_shape is not None:
        new_para = (
            "The next step is orchestration — a governed way for Sourcing Agent "
            "to run across several approved environments at once. Scheduled, "
            "supervised, and fully auditable."
        )
        set_text_preserve_formatting(p2_shape.text_frame, new_para)

    # --- Slide 9: single-line title, remove SUMMARY eyebrow, shrink title size ---
    s9 = slides[8]
    # Remove the SUMMARY eyebrow (Text 0).
    eyebrow = find_shape_by_text(s9, "SUMMARY")
    if eyebrow is not None:
        remove_shape(eyebrow)
    # Rewrite title to a single line and reduce the font size.
    title = find_shape_by_text(s9, "Built in-house")
    if title is not None:
        tf = title.text_frame
        set_text_preserve_formatting(tf, "Built in-house. Running today.")
        # Shrink the font so the close is visibly quieter than the cover.
        # Cover title is the default large size; aim for ~60pt here.
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(60)

    # --- Save ---
    prs.save(str(DST))
    print(f"wrote {DST}")


if __name__ == "__main__":
    main()
